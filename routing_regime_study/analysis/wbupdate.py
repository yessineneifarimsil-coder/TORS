"""Add regime/decision sheets to the workbook. Existing sheets are untouched."""
import openpyxl, json, csv, statistics as st
from openpyxl.styles import Font
from collections import defaultdict
SRC="/root/.claude/uploads/35501cb2-bbe1-5745-801e-da4fe3f5e9fd/144859a5-Routing_Strategy_Selection_Corrected.xlsx"
DST="paper/Routing_Strategy_Selection_Regime.xlsx"
O=json.load(open("results.json"))
wb=openpyxl.load_workbook(SRC)
before=[ws.title for ws in wb.worksheets]
B=Font(bold=True)
def newsheet(name,title,note):
    ws=wb.create_sheet(name); ws["A2"]=title; ws["A2"].font=Font(bold=True,size=12)
    ws["A4"]=note; return ws
# --- REGIME_TRANSITION ---
ws=newsheet("REGIME_TRANSITION","Regime structure of routing-policy preference",
  "Switching margin m = C(SP) - C(DTT) in common cost units, balanced profile. Negative: SP preferred. Derived from the 432 completed runs; no new simulation.")
hdr=["Signal","Restriction","m @ q=720","m @ q=1200","m @ q=1680","Crossing bracket (veh/h)"]
for j,h in enumerate(hdr,1): ws.cell(7,j,h).font=B
r=8
for k,v in O['transition'].items():
    sig,res=k.split("|")
    ws.cell(r,1,sig); ws.cell(r,2,res)
    for j,m in enumerate(v['margins']): ws.cell(r,3+j,m)
    ws.cell(r,6,f"({v['bracket'][0]}, {v['bracket'][1]}]" if v['bracket'] else "no sign change")
    r+=1
r+=2
ws.cell(r,1,"Worst-policy regret by demand").font=B; r+=1
for q,val in O['worst_policy_regret'].items(): ws.cell(r,1,f"q={q}"); ws.cell(r,2,val); r+=1
r+=1
ws.cell(r,1,"Seed stability").font=B; r+=1
ws.cell(r,1,"contexts with identical winner in all 3 seeds"); ws.cell(r,2,O['seed_stable_winner']); ws.cell(r,3,"of 24"); r+=1
ws.cell(r,1,"contexts with stable margin sign"); ws.cell(r,2,O['seed_stable_margin_sign']); ws.cell(r,3,"of 24"); r+=1
ws.cell(r,1,"unstable contexts"); ws.cell(r,2,", ".join(O['seed_unstable_ctx'])); r+=1
ws.cell(r,1,"LIMITATION"); ws.cell(r,2,"Demand grid is 480 veh/h coarse: the reversal is BRACKETED, not localised. No threshold estimate is made.")
# --- DECISION_EVALUATION ---
ws=newsheet("DECISION_EVALUATION","Held-out routing-policy selection (leave-one-context-out)",
  "Portfolio SP/DTT/TECO10. Model and reference scales refit per fold; test context never seen. No new simulation.")
for j,h in enumerate(["Method","Mean regret","Max regret","Optimal /24","Within 1% /24"],1): ws.cell(7,j,h).font=B
r=8
rows=[("Fixed SP",O['fixed']['SP']),("Fixed DTT (best fixed)",O['fixed']['DTT']),
      ("Fixed TECO10",O['fixed']['TECO10']),("Adaptive selector (held-out)",O['adaptive_loco']),
      ("Hindsight Best-Policy Benchmark",O['hindsight'])]
for lab,d in rows:
    ws.cell(r,1,lab); ws.cell(r,2,d['mean']); ws.cell(r,3,d['max']); ws.cell(r,4,d['zero']); ws.cell(r,5,d['near']); r+=1
r+=1
for lab,val in [("Available headroom (fixed DTT -> hindsight)",O['headroom_mean']),
                ("Headroom captured (%)",O['captured_pct']),
                ("Prediction MAE time (s)",O['pred_mae_time']),("Prediction RMSE time (s)",O['pred_rmse_time']),
                ("Prediction MAE CO2 (g)",O['pred_mae_co2']),("Prediction RMSE CO2 (g)",O['pred_rmse_co2'])]:
    ws.cell(r,1,lab).font=B; ws.cell(r,2,val); r+=1
r+=2
ws.cell(r,1,"EXTRAPOLATION STRESS TEST (hold out an entire demand regime)").font=B; r+=1
for j,h in enumerate(["Held-out q","Selector mean regret","Fixed DTT mean regret","Selector chose"],1): ws.cell(r,j,h).font=B
r+=1
for q in ("720","1200","1680"):
    e=O['extrapolation'][q]
    ws.cell(r,1,int(q)); ws.cell(r,2,e['cart']); ws.cell(r,3,e['dtt']); ws.cell(r,4,", ".join(e['chosen'])); r+=1
ws.cell(r,1,"POOLED").font=B; ws.cell(r,2,O['extrap_pooled_cart']); ws.cell(r,3,O['extrap_pooled_dtt'])
ws.cell(r,4,f"captured {O['extrap_captured_pct']:.1f}% -> selector WORSE than fixed"); r+=2
ws.cell(r,1,"Per-context held-out decisions").font=B; r+=1
for j,h in enumerate(["Context","Selected policy"],1): ws.cell(r,j,h).font=B
r+=1
for c,p in O['chosen_loco'].items(): ws.cell(r,1,c); ws.cell(r,2,p); r+=1
# --- NEGATIVE_CONTROLS ---
ws=newsheet("NEGATIVE_CONTROLS","Audited negative controls",
  "Retained, not deleted. These delimit what the positive result means.")
for j,h in enumerate(["Control","Observed value","Interpretation"],1): ws.cell(7,j,h).font=B
r=8
for a,b,c in [
 ("ECO outcome-identical to SP",f"{O['eco_identical_to_sp']}/24 contexts","No independent environmental decision dimension; ECO removed from portfolio"),
 ("Observer route-order selection loss","21.62%","Fails declared 10% screen despite 14.71% aggregate WAPE"),
 ("Observer best-route hits","3/8 panels","Good average prediction did not confer correct ranking"),
 ("Selector decisions across w_time 0..1","identical in all 24 contexts","Preference layer is near-inert; regime pattern is NOT a weight artefact"),
 ("Hindsight label change across weights","1/24 contexts (D011)","Only at the pure-emissions extreme"),
 ("Insertion-boundary winner flips","0/24 contexts","Regime structure not an artefact of the emission accounting boundary"),
 ("Pilot adaptive headroom","0 (DTT best in 30/30)","Earlier environment had no regime reversal"),
]:
    ws.cell(r,1,a); ws.cell(r,2,b); ws.cell(r,3,c); r+=1
for s in ("REGIME_TRANSITION","DECISION_EVALUATION","NEGATIVE_CONTROLS"):
    w=wb[s]
    for col,wd in zip("ABCDEF",[42,24,22,18,18,26]): w.column_dimensions[col].width=wd
wb.save(DST)
print("original sheets:",len(before)," -> new file sheets:",len(wb.worksheets))
print("added:",[t for t in [ws.title for ws in wb.worksheets] if t not in before])
