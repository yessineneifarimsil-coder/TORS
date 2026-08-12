from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# PROJECT PATHS
# ============================================================

ROOT = Path(__file__).resolve().parents[1]

DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(exist_ok=True)

OUTPUT_FILE = DATA_DIR / "SWFC_research_workbook.xlsx"


# ============================================================
# CREATE WORKBOOK
# ============================================================

wb = Workbook()


# ============================================================
# SHEET 1 — README
# ============================================================

ws = wb.active
ws.title = "README"

readme = [
    ["SWFC Research Workbook", ""],
    ["Purpose",
     "Central workbook for documenting the experimental design, "
     "criteria, alternatives, synthetic data, SHAP weights and MCDM results."],
    ["Status",
     "Research development — synthetic benchmark stage."],
    ["Important",
     "Synthetic values must not be presented as measured Tunisian observations."],
    ["Expert data",
     "No expert judgments are included unless genuine expert elicitation is performed."],
    ["Computational engine",
     "Python"],
]

for row in readme:
    ws.append(row)


# ============================================================
# SHEET 2 — CRITERIA
# ============================================================

ws = wb.create_sheet("Criteria")

headers = [
    "ID",
    "Criterion",
    "Dimension",
    "Type",
    "Definition",
    "Unit",
    "Generation method",
    "Source",
]

ws.append(headers)

criteria = [
    ["C1", "Congestion", "Mobility", "Cost",
     "Traffic congestion / delay exposure",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C2", "Travel-time variability", "Mobility", "Cost",
     "Variability of travel time",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C3", "Capacity utilisation", "Mobility", "Cost",
     "Degree of capacity saturation",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C4", "Accident exposure", "Safety", "Cost",
     "Exposure to traffic safety risk",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C5", "Incident response", "Safety", "Benefit",
     "Expected improvement in incident response",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C6", "CO2 intensity", "Environment", "Cost",
     "Traffic-related CO2 emissions intensity",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C7", "Fuel consumption", "Environment", "Cost",
     "Traffic-related fuel consumption",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C8", "Noise", "Environment", "Cost",
     "Traffic noise exposure",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C9", "Deployment cost", "Implementation", "Cost",
     "Cost associated with ITS deployment",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C10", "Infrastructure readiness", "Implementation", "Benefit",
     "Readiness of existing infrastructure",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C11", "Network extensibility", "Implementation", "Benefit",
     "Ability to scale to a wider network",
     "TBD", "TBD", "Synthetic benchmark"],
    
    ["C12", "Public-transport satisfaction", "Social", "Benefit",
     "Expected public-transport/user benefit",
     "TBD", "TBD", "Synthetic benchmark"],
]

for row in criteria:
    ws.append(row)


# ============================================================
# SHEET 3 — ALTERNATIVES
# ============================================================

ws = wb.create_sheet("Alternatives")

ws.append(["ID", "Alternative", "Description"])

alternatives = [
    ["A1", "Adaptive Traffic Signal Control (ATSC)",
     "Adaptive traffic signal control"],
    
    ["A2", "Variable Message Signs (VMS)",
     "Dynamic traveller information through variable message signs"],
    
    ["A3", "Smart Traveler Information Systems (STIS)",
     "Traveler information and decision-support services"],
    
    ["A4", "V2I / V2X Communication",
     "Vehicle-to-infrastructure / vehicle-to-everything communication"],
    
    ["A5", "Eco-Routing Systems",
     "Routing strategies targeting energy and emissions reduction"],
    
    ["A6", "Integrated Traffic Management Control Center (TMCC)",
     "Integrated traffic monitoring and management"],
]

for row in alternatives:
    ws.append(row)


# ============================================================
# SHEET 4 — FUZZY SCALE
# ============================================================

ws = wb.create_sheet("Fuzzy_Scale")

ws.append([
    "Code",
    "Linguistic term",
    "Lower",
    "Middle",
    "Upper"
])

fuzzy_scale = [
    ["VL", "Very Low", 0.00, 0.00, 0.10],
    ["L",  "Low", 0.10, 0.20, 0.30],
    ["ML", "Medium-Low", 0.20, 0.35, 0.50],
    ["M",  "Medium", 0.35, 0.50, 0.65],
    ["MH", "Medium-High", 0.50, 0.65, 0.80],
    ["H",  "High", 0.70, 0.80, 0.90],
    ["VH", "Very High", 0.85, 0.95, 1.00],
]

for row in fuzzy_scale:
    ws.append(row)


# ============================================================
# SHEET 5 — EXPERIMENT DESIGN
# ============================================================

ws = wb.create_sheet("Experiment_Design")

ws.append([
    "Parameter",
    "Value",
    "Status",
    "Notes"
])

experiment_parameters = [
    ["Number of corridors", "TBD", "To decide",
     "Must be justified by experimental design."],

    ["Scenarios per corridor", "TBD", "To decide",
     "Must provide sufficient variation."],

    ["Validation method", "Grouped cross-validation",
     "Proposed", "Groups = corridor."],

    ["Primary model", "XGBoost Regressor",
     "Proposed", "Primary predictive model."],

    ["Primary explanation", "SHAP",
     "Proposed", "Global mean absolute SHAP importance."],

    ["MCDM method", "Fuzzy CoCoSo",
     "Proposed", "Primary ranking method."],

    ["Sensitivity", "Criterion removal",
     "Proposed", "One criterion removed at a time."],

    ["Robustness", "Weight perturbation",
     "Proposed", "Weights renormalized after perturbation."],

    ["Synthetic data", "YES",
     "Current stage", "Not empirical Tunisian measurements."],
]

for row in experiment_parameters:
    ws.append(row)


# ============================================================
# SHEET 6 — CORRIDOR DATA
# ============================================================

ws = wb.create_sheet("Corridor_Data")

ws.append(
    ["corridor_id", "scenario_id"]
    + [f"C{i}" for i in range(1, 13)]
    + ["target_Y", "source", "notes"]
)


# ============================================================
# SHEET 7 — FUZZY MATRIX
# ============================================================

ws = wb.create_sheet("Fuzzy_Matrix")

ws.append([
    "alternative_id",
    "criterion_id",
    "linguistic_code",
    "lower",
    "middle",
    "upper",
    "source",
    "notes"
])


# ============================================================
# SHEET 8 — SHAP WEIGHTS
# ============================================================

ws = wb.create_sheet("SHAP_Weights")

ws.append([
    "criterion_id",
    "criterion",
    "mean_abs_shap",
    "normalized_weight",
    "rank",
    "seed",
    "validation_scheme",
    "model"
])


# ============================================================
# SHEET 9 — MCDM RESULTS
# ============================================================

ws = wb.create_sheet("MCDM_Results")

ws.append([
    "alternative_id",
    "alternative",
    "method",
    "score",
    "rank",
    "weight_source"
])


# ============================================================
# SHEET 10 — SENSITIVITY
# ============================================================

ws = wb.create_sheet("Sensitivity")

ws.append([
    "experiment",
    "parameter",
    "scenario",
    "score",
    "rank",
    "spearman_rho",
    "notes"
])


# ============================================================
# FORMATTING
# ============================================================

for ws in wb.worksheets:

    ws.freeze_panes = "A2"

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for column in ws.columns:

        max_length = 0

        for cell in column:

            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        width = min(max_length + 3, 50)

        ws.column_dimensions[
            column[0].column_letter
        ].width = width


# ============================================================
# SAVE
# ============================================================

wb.save(OUTPUT_FILE)

print()
print("=" * 60)
print("SWFC RESEARCH WORKBOOK CREATED")
print("=" * 60)
print(f"Location: {OUTPUT_FILE}")
print()